"""Simple Word report: app URL, environment, username, password, Jira card, description,
then Step 1..n each with its screenshot. Zero AI tokens."""

from __future__ import annotations

import getpass
import json
import os
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocxDocument
from docx.shared import Inches, Pt, RGBColor

from .config import QA, app_env, fail, load_env, rel, remember, state

SECRET_KEY = re.compile(r'"(password|passwd|pwd|pass|secret|token|apiKey)"\s*:', re.I)
COLOR = {"Pass": RGBColor(0x1E, 0x7B, 0x34), "Fail": RGBColor(0xC0, 0x00, 0x00),
         "Blocked": RGBColor(0xB2, 0x6B, 0x00), "Not Run": RGBColor(0x66, 0x66, 0x66)}


def _load(test_id: str) -> dict[str, Any]:
    file = QA / "testcases" / f"{test_id}.json"
    if not file.exists():
        fail(f"Not found: {rel(file)}")
    raw = file.read_text(encoding="utf-8")
    if SECRET_KEY.search(raw):
        fail(f"{file.name} has a password/secret field. Passwords belong in qa/.env only.")
    tc: dict[str, Any] = json.loads(raw)
    missing = [k for k in ("id", "title", "environment", "steps") if not tc.get(k)]
    if missing:
        fail(f"{file.name} is missing: {', '.join(missing)}")
    return tc


def status_of(tc: dict[str, Any]) -> str:
    if tc.get("status"):
        return str(tc["status"])
    results = [s.get("result") for s in tc["steps"]]
    if "Fail" in results:
        return "Fail"
    if "Blocked" in results:
        return "Blocked"
    return "Pass" if all(r == "Pass" for r in results) else "Not Run"


def _description(tc: dict[str, Any]) -> str | None:
    if tc.get("description"):
        return str(tc["description"])
    md = QA / "jira" / f"{tc['id']}.md"
    if not md.exists():
        return None
    part = md.read_text(encoding="utf-8").split("## Description", 1)
    text = part[1].split("\n## ", 1)[0] if len(part) > 1 else ""
    return " ".join(text.split())[:600] or None


def _line(doc: DocxDocument, label: str, value: str | None) -> None:
    if not value:
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.add_run(f"{label}: ").bold = True
    p.add_run(value)


def _result(p: Any, result: str | None) -> None:
    if result:
        run = p.add_run(result)
        run.bold = True
        run.font.color.rgb = COLOR.get(result, COLOR["Not Run"])


def _section(doc: DocxDocument, tc: dict[str, Any], first: bool, hide_password: bool) -> None:
    env = tc["environment"]
    shots = QA / "evidence" / tc["id"]
    taken = sorted(shots.glob("step-*.png"))
    when = datetime.fromtimestamp(taken[-1].stat().st_mtime) if taken else datetime.now()

    app = env.get("app")
    user = (tc.get("login") or {}).get("user") or (app_env(app, "USER") if app else None)
    password = ("(hidden)" if hide_password else app_env(app, "PASS")) if app and user else None
    env_text = ", ".join(x for x in (env.get("name"), env.get("build") and f"build {env['build']}", env.get("browser")) if x)

    heading = doc.add_heading(f"{tc['id']}: {tc['title']}", level=1)
    heading.paragraph_format.page_break_before = not first
    tester = os.environ.get("QA_TESTER") or getpass.getuser()
    p = doc.add_paragraph(f"Date: {when:%Y-%m-%d %H:%M}    Tester: {tester}    Result: ")
    _result(p, status_of(tc))
    _line(doc, "Application URL", env.get("url"))
    _line(doc, "Environment", env_text)
    _line(doc, "Username", user)
    _line(doc, "Password", password)
    _line(doc, "Jira card", tc.get("jira"))
    _line(doc, "Description", _description(tc))

    for n, step in enumerate(tc["steps"], start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.keep_with_next = True
        p.add_run(f"Step {n}: ").bold = True
        p.add_run(f"{step['action']}  ")
        _result(p, step.get("result"))
        if step.get("result") not in (None, "Pass") and step.get("actual"):
            _line(doc, "Actual", step["actual"])
        shot = shots / f"step-{n:02d}.png"
        if shot.exists():
            doc.add_picture(str(shot), width=Inches(6.6))


def build(ids: list[str], hide_password: bool = False, all_cases: bool = False) -> None:
    """No ids: the last test you worked on (or every test case with --all)."""
    load_env()
    folder = QA / "testcases"
    last = state().get("last_id")
    if all_cases:
        ids = sorted(p.stem for p in folder.glob("*-*.json"))
    elif not ids and last:
        ids = [last]
    if not ids:
        fail("Which test? Give an id, e.g. lct report PROJ-123 (or --all)")
    cases = [_load(i) for i in ids]

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(0.7)
    doc.core_properties.title = ", ".join(tc["id"] for tc in cases)
    for i, tc in enumerate(cases):
        _section(doc, tc, i == 0, hide_password)

    out_dir: Path = QA / "reports"
    out_dir.mkdir(parents=True, exist_ok=True)
    name = f"{cases[0]['id'] if len(cases) == 1 else 'Test'}_Report_{datetime.now():%Y-%m-%d-%H-%M}.docx"
    try:
        doc.save(str(out_dir / name))
    except PermissionError:
        fail("Close the report in Word and run again.")
    remember(report=rel(out_dir / name))
    print(f"{rel(out_dir / name)} | " + ", ".join(f"{tc['id']} {status_of(tc)}" for tc in cases))
